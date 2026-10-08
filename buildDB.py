"""buildDB.py — Befüllt die semantische Kontextdatenbank aus PDFs und DOCX.

Liest BSI-Bausteine, Guidelines und interne Vorgaben aus PDF- oder
Word-Dateien ein, zerlegt sie in Chunks, erzeugt lokale Embeddings
(TF-IDF + SVD, kein Modell-Download) und speichert alles persistent
in ``context_db/``.

Aufruf:
    python buildDB.py <datei-oder-ordner> [--rebuild] [--db-dir context_db] [-v]
    python buildDB.py --stats [--db-dir context_db]

Beispiele:
    python buildDB.py docs/bsi/                     # Ordner mit PDFs/DOCX einlesen
    python buildDB.py richtlinie.docx --rebuild     # Einzeldatei, DB neu bauen
    python buildDB.py docs/bsi/ --stats             # Statistik anzeigen
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from context import ContextDB, chunk_text


def extract_pdf_text(path: Path) -> str:
    """Extrahiert Text aus einem PDF (alle Seiten, mit Seitenmarkern)."""
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text() or ""
            parts.append(f"--- Seite {i + 1} ---\n{page_text}")
    return "\n".join(parts)


def extract_docx_text(path: Path) -> str:
    """Extrahiert Text aus einem DOCX (Absätze, plus Tabellen).

    Überschriften (Heading-Styles) bleiben als eigene Absätze erhalten —
    ``chunk_text`` erkennt sie an der BSI-Überschriften-Heuristik.
    """
    import docx

    document = docx.Document(str(path))
    parts = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def collect_files(source: Path) -> list[Path]:
    """Sammelt alle PDF/DOCX-Dateien aus einer Datei oder einem Ordner."""
    if source.is_file():
        return [source]
    files: list[Path] = []
    for pattern in ("*.pdf", "*.docx", "*.PDF", "*.DOCX"):
        files.extend(source.rglob(pattern))
    return sorted(set(files))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Semantische Kontextdatenbank aus PDFs/DOCX befüllen"
    )
    parser.add_argument("source", nargs="?", help="Datei oder Ordner mit PDFs/DOCX")
    parser.add_argument(
        "--rebuild",
        action="store_true",
        help="Bestehende DB verwerfen und komplett neu bauen",
    )
    parser.add_argument(
        "--db-dir",
        default="context_db",
        help="Speicherort der DB (default: context_db/)",
    )
    parser.add_argument(
        "--stats", action="store_true", help="Nur Statistiken anzeigen und beenden"
    )
    parser.add_argument("-v", "--verbose", action="store_true", help="Mehr Ausgabe")
    args = parser.parse_args()

    db = ContextDB(persist_dir=args.db_dir)

    if args.stats:
        if not db.load():
            print(f"Keine Datenbank in {args.db_dir}/ gefunden.")
            return 1
        stats = db.stats()
        print(f"Chunks:    {stats['chunks']}")
        print(f"Dimension: {stats['dims']}")
        print(f"Quellen:   {len(stats['sources'])}")
        for s in stats["sources"]:
            print(f"  - {s}")
        return 0

    if not args.source:
        parser.error("Quelle (Datei/Ordner) fehlt — oder --stats verwenden.")

    source = Path(args.source)
    if not source.exists():
        print(f"FEHLER: {source} existiert nicht.")
        return 1

    if args.rebuild:
        print("Rebuild: bestehende DB wird verworfen.")
        db.clear()
    else:
        db.load()  # bestehende DB laden, falls vorhanden

    files = collect_files(source)
    if not files:
        print(f"Keine PDF/DOCX-Dateien in {source} gefunden.")
        return 1

    all_chunks = []
    for path in files:
        print(f"Lese: {path.name}")
        try:
            if path.suffix.lower() == ".pdf":
                text = extract_pdf_text(path)
            else:
                text = extract_docx_text(path)
        except Exception as exc:
            print(f"  FEHLER: {exc}")
            continue
        chunks = chunk_text(text, source=path.name)
        all_chunks.extend(chunks)
        print(f"  -> {len(chunks)} Chunks, {len(text)} Zeichen")
        if args.verbose:
            for c in chunks[:3]:
                preview = c.text[:80].replace("\n", " ")
                print(f"     [{c.chunk_index}] {preview}...")

    if all_chunks:
        added = db.add_chunks(all_chunks)
        print(f"\n{added} neue Chunks hinzugefügt (einmaliges Retraining).")

    db.save()
    stats = db.stats()
    print(f"\nFertig: {stats['chunks']} Chunks, {stats['dims']} Dimensionen.")
    print(f"Gespeichert in: {stats['persist_dir']}/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
